#!/usr/bin/env python3
"""
Đối chiếu log mô phỏng của testbench S5 với các item trong vplan.

Dùng:
  python3 check_sim_log.py <log> <vplan.md|.xlsx|.csv|.docx> [--id-re REGEX]
  python3 check_sim_log.py --items <vplan> [--id-re REGEX]

Định dạng log (default_tb_rules.md):
  -- <item name> test --                                       đầu mỗi item vplan (tên đọc được)
  [<time>] <case> PASS | expected: <exp> | actual: <act>       mỗi lần checker so sánh
  [<time>] <case> FAIL | expected: <exp> | actual: <act>
  [SUMMARY] PASS=<n> FAIL=<n>                                  cuối mô phỏng
Case thuộc item của header gần nhất phía trên.

Item trong vplan = (ID, tên):
  --id-re mặc định 'TC_[A-Z0-9_]+_\\d{3}' (regex có nhóm bắt thì lấy nhóm 1); tên = phần sau ID trên
  cùng dòng (heading `### TC_X_001 – <tên>`), hoặc ô kế tiếp nếu dòng là hàng bảng.
  Vplan không có ID dạng TC_… → bảng có cột đầu "ID" (vd. template checklist 1, 2, 3); tên = cột cụ thể
  nhất trong Item / Sub item 1 / Sub item 2 (thường Sub item 2), chỉ thêm cột phía trước (nối " / ")
  khi cần để phân biệt với item khác; bỏ ô trống, "—". Không có các cột đó thì cột thứ hai.
  --items: in danh sách ID → header mong đợi rồi thoát (để viết đúng tên trong testbench).
Header khớp item theo tên (không phân biệt hoa thường, khoảng trắng); header ghi đúng ID cũng được nhận
(dùng khi vplan không có tên hoặc tên trùng).

Kiểm: item vplan chưa chạy / không có case · case FAIL · dòng case sai định dạng · case nằm trước mọi
header · header không có trong vplan / trùng · thiếu, trùng hoặc lệch [SUMMARY] · [TIMEOUT]/$fatal/FATAL.
Mã thoát 0 = mọi item trong vplan đều PASS và log nhất quán; 1 = có vấn đề.
"""
import argparse
import csv
import re
import sys
from pathlib import Path

DEFAULT_ID_RE = r"TC_[A-Z0-9_]+_\d{3}"
HEADER = re.compile(r"^--\s+(?P<name>.+?)\s+test\s+--\s*$")
CASE = re.compile(r"^\[(?P<t>\d+)\]\s+(?P<case>\S+)\s+(?P<res>PASS|FAIL)\s+\|\s+expected:\s*(?P<exp>.*?)"
                  r"\s+\|\s+actual:\s*(?P<act>.*?)\s*$")
CASE_LIKE = re.compile(r"^\[\d+\]\s+\S+\s+(PASS|FAIL)\b")
SUMMARY = re.compile(r"\[SUMMARY\]\s+PASS=(\d+)\s+FAIL=(\d+)")
EMPTY = {"", "-", "–", "—"}


def clean(s):
    return re.sub(r"[`*]", "", s).strip()


def row(cells):
    return "| " + " | ".join("" if c is None else str(c).replace("\n", " ") for c in cells) + " |"


def vplan_lines(p: Path):
    """Dòng văn bản của vplan; hàng bảng (md/csv/xlsx/docx) đều ở dạng '| a | b |'."""
    ext = p.suffix.lower()
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        return [row(r) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
    if ext == ".docx":
        import docx
        d = docx.Document(p)
        out = [para.text for para in d.paragraphs]
        for t in d.tables:
            out += [row(c.text for c in r.cells) for r in t.rows]
        return out
    text = p.read_text(encoding="utf-8", errors="replace")
    if ext == ".csv":
        return [row(r) for r in csv.reader(text.splitlines())]
    out, fence = [], False
    for line in text.splitlines():          # bỏ khối ``` (ví dụ/mẫu trong vplan)
        if line.strip().startswith("```"):
            fence = not fence
        elif not fence:
            out.append(line)
    return out


def cells_of(line):
    return [clean(c) for c in line.strip().strip("|").split("|")]


def items_by_id(lines, id_re):
    rx = re.compile(id_re)
    items = {}
    for line in lines:
        m = rx.search(line)
        if not m:
            continue
        tid = m.group(1) if rx.groups else m.group(0)
        if tid in items:
            continue
        rest = line[m.end():]
        if line.strip().startswith("|"):
            name = next((c for c in cells_of(rest) if c not in EMPTY), "")
        else:
            name = clean(re.sub(r"^[\s\-–—:.)\]]+", "", rest))
        items[tid] = name
    return items


def items_by_table(lines):
    """Bảng có cột đầu 'ID' – dùng khi vplan không có ID dạng TC_…"""
    rows, header, first = {}, None, True
    for line in lines:
        if not line.strip().startswith("|"):
            header, first = None, True
            continue
        c = cells_of(line)
        if first:                           # chỉ hàng đầu của bảng là tiêu đề
            first = False
            header = [h.lower() for h in c] if c and c[0].lower() == "id" else None
            continue
        if header is None:
            continue
        if set("".join(c)) <= set("-: ") or not c[0]:
            continue
        cols = [i for i, h in enumerate(header) if h.startswith(("item", "sub item"))] or [1]
        rows.setdefault(c[0], [c[i] for i in cols if i < len(c) and c[i] not in EMPTY])
    return {tid: short_name(f, rows.values()) for tid, f in rows.items()}


def short_name(fields, all_fields):
    """Ít cột nhất (tính từ cột cụ thể nhất) đủ để tên khác mọi item khác."""
    def tail(f, k):
        return norm(" / ".join(f[-k:]))
    for k in range(1, len(fields) + 1):
        if sum(tail(f, k) == tail(fields, k) for f in all_fields) == 1:
            return " / ".join(fields[-k:])
    return " / ".join(fields)


def norm(s):
    s = re.sub(r"\s+", " ", s.strip().lower())
    return re.sub(r" test$", "", s)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", metavar="<log> <vplan> | --items <vplan>")
    ap.add_argument("--id-re", default=DEFAULT_ID_RE)
    ap.add_argument("--items", action="store_true", help="in ID → header mong đợi rồi thoát")
    a = ap.parse_args()
    if len(a.files) != (1 if a.items else 2):
        ap.error("cần <log> <vplan>, hoặc --items <vplan>")
    log, vplan = (None, a.files[0]) if a.items else a.files

    vlines = vplan_lines(Path(vplan))
    items, how = items_by_id(vlines, a.id_re), f"--id-re {a.id_re}"
    if not items and a.id_re == DEFAULT_ID_RE:
        items, how = items_by_table(vlines), "cột ID của bảng vplan"
    if a.items:
        print(f"Vplan: {len(items)} item ({how})")
        for tid, name in items.items():
            print(f"  {tid:<20} -- {name or tid} test --")
        sys.exit(0 if items else 1)
    problems = []
    if not items:
        problems.append("không tìm thấy item nào trong vplan – chỉ định --id-re theo quy ước ID của vplan")

    # header (tên item hoặc ID) → ID
    by_name = {}
    for tid, name in items.items():
        if name:
            by_name.setdefault(norm(name), []).append(tid)
    dup = {k: v for k, v in by_name.items() if len(v) > 1}
    for k, v in dup.items():
        problems.append(f"tên item '{k}' trùng trong vplan ({', '.join(v)}) – header các item này phải ghi ID")
    lookup = {k: v[0] for k, v in by_name.items() if len(v) == 1}
    lookup.update({norm(t): t for t in items})

    lines = Path(log).read_text(encoding="utf-8", errors="replace").splitlines()
    results, extra = {}, []
    cur, n_pass, n_fail, summaries = None, 0, 0, []
    for n, line in enumerate(lines, 1):
        s = line.strip()
        m = HEADER.match(s)
        if m:
            tid = lookup.get(norm(m.group("name")))
            cur = tid or f"?{m.group('name')}"
            if tid is None:
                extra.append(m.group("name"))
            if cur in results:
                problems.append(f"L{n}: header item '{m.group('name')}' xuất hiện lần nữa")
            results.setdefault(cur, [])
            continue
        m = CASE.match(s)
        if m:
            if m.group("res") == "PASS":
                n_pass += 1
            else:
                n_fail += 1
            if cur is None:
                problems.append(f"L{n}: case '{m.group('case')}' nằm trước mọi header '-- <item name> test --'")
            else:
                results[cur].append((n, m))
            continue
        if CASE_LIKE.match(s):
            problems.append(f"L{n}: dòng case sai định dạng "
                            f"(cần '[<time>] <case> PASS|FAIL | expected: <exp> | actual: <act>'): {s[:120]}")
            continue
        m = SUMMARY.search(s)
        if m:
            summaries.append((int(m.group(1)), int(m.group(2))))
        if re.search(r"\[TIMEOUT\]|\$fatal|FATAL", s):
            problems.append(f"L{n}: mô phỏng dừng bất thường: {s[:120]}")

    n_run = sum(1 for k in results if not k.startswith("?"))
    print(f"Vplan: {len(items)} item ({how}) · Log: {n_run} item, {n_pass} case PASS, {n_fail} case FAIL")

    bad = 0
    for tid, name in items.items():
        cases = results.get(tid)
        if cases is None:
            st = f"KHÔNG CHẠY (không có header '-- {name or tid} test --')"
        elif not cases:
            st = "KHÔNG CÓ KẾT QUẢ (header không có case nào)"
        else:
            fails = [(n, m) for n, m in cases if m.group("res") == "FAIL"]
            st = f"FAIL ({len(fails)}/{len(cases)} case)" if fails else f"PASS ({len(cases)} case)"
        bad += not st.startswith("PASS")
        print(f"  {tid:<20} {name[:40]:<40} {st}")
        for n, m in (cases or []):
            if m.group("res") == "FAIL":
                print(f"      L{n} [{m.group('t')}] {m.group('case')}: "
                      f"expected {m.group('exp')} · actual {m.group('act')}")
    for name in extra:
        print(f"  {'?':<20} {name[:40]:<40} THỪA (header không khớp tên/ID item nào trong vplan)")
        bad += 1

    if not summaries:
        problems.append("thiếu dòng [SUMMARY] PASS=<n> FAIL=<n>")
    else:
        sp, sf = summaries[-1]
        if (sp, sf) != (n_pass, n_fail):
            problems.append(f"[SUMMARY] PASS={sp} FAIL={sf} lệch với số dòng case trong log "
                            f"(PASS={n_pass} FAIL={n_fail})")
        if len(summaries) > 1:
            problems.append(f"có {len(summaries)} dòng [SUMMARY]")
    for p in problems:
        print("  [VẤN ĐỀ] " + p)

    ok = bad == 0 and not problems
    print("KẾT QUẢ:", "TẤT CẢ PASS" if ok else "CÓ VẤN ĐỀ")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
